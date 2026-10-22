#!/usr/bin/env python3
"""Review sheets for native speakers, and reading their corrections back.

    python3 tools/translations/sheet.py build texts.json OUT_DIR
        -> one .xlsx per language (UZ, KG): instructions + every text, with the picture for words
    python3 tools/translations/sheet.py read FILLED.xlsx corrections.json
        -> the corrections, ready for: node tools/translations/strings.js apply corrections.json

Needs openpyxl and Pillow (pip install openpyxl pillow).
"""
import io
import json
import os
import re
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..'))
IMAGES = os.path.join(ROOT, 'backend', 'images')

LANG_NAMES = {'UZ': ('Узбекский', 'узбекском'), 'KG': ('Кыргызский', 'кыргызском')}
FONT = 'Arial'
HEADERS = ['№', 'Картинка', 'Где в игре', 'Кто видит', 'По-русски (смысл)', 'Текст сейчас', 'Верно?', 'Как правильно', 'Комментарий', 'id']
WIDTHS = [6, 10, 26, 11, 34, 38, 9, 38, 28, 8]
FILL_IN = PatternFill('solid', start_color='FFF6C4')
HEADER_FILL = PatternFill('solid', start_color='3A6FF0')
EXAMPLE_FILL = PatternFill('solid', start_color='EEF3FF')
THIN = Side(style='thin', color='D3DCF2')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical='center')
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)


def picture_file(key):
    if not key:
        return None
    key = re.sub(r'\d+$', '', key)
    path = os.path.join(IMAGES, f'{key}.webp')
    return path if os.path.exists(path) else None


def thumbnail(path, size=52):
    from PIL import Image
    with Image.open(path) as source:
        image = source.convert('RGBA')
        image.thumbnail((size, size))
        canvas = Image.new('RGBA', (size, size), (255, 255, 255, 0))
        canvas.paste(image, ((size - image.width) // 2, (size - image.height) // 2), image)
        data = io.BytesIO()
        canvas.save(data, format='PNG')
        data.seek(0)
        return data


def context_of(item):
    if item['screen'] == 'Названия букв':
        letter = item['id'].rsplit('.', 1)[-1]
        return f'название буквы «{letter}» (так её произносит Ирби)'
    if item['kind'] == 'syllables':
        return (item['ru'] + ' — ' if item['ru'] else '') + 'слово по слогам, слоги через дефис'
    return item['ru']


def write_instructions(sheet, lang):
    name, where = LANG_NAMES[lang]
    sheet.title = 'Как заполнять'
    sheet.column_dimensions['A'].width = 4
    for column, width in zip('BCDEFGHI', [24, 30, 36, 10, 34, 26, 4, 4]):
        sheet.column_dimensions[column].width = width
    sheet['B2'] = f'Letter Loom: проверка текстов на {where} языке'
    sheet['B2'].font = Font(name=FONT, size=16, bold=True)
    lines = [
        'Это все фразы и слова, которые игра показывает и произносит. Их писал не носитель языка, поэтому в них могут быть ошибки.',
        f'Откройте лист «{name}». В столбце «Текст сейчас» — как сейчас в игре, рядом — что это значит по-русски и картинка для слов.',
        'Если всё верно — ничего не пишите (можно поставить «Да» в столбце «Верно?»).',
        'Если есть ошибка или так не говорят с детьми — напишите правильный вариант в столбце «Как правильно».',
        'Заполнять можно только жёлтые столбцы: «Верно?», «Как правильно», «Комментарий».',
        'Части в фигурных скобках, например {n} или {name}, не меняйте: туда игра сама подставит число или имя.',
        'В словах по слогам ставьте дефис между слогами: ' + ('AL-MA.' if lang == 'UZ' else 'АЛ-МА.'),
        'Если Ирби неправильно произносит слово (ударение, звук) — напишите об этом в «Комментарий».',
        'Сначала идут тексты для детей, в конце — тексты для взрослых (режим учителя).',
        'Когда закончите, сохраните файл и отправьте его обратно.'
    ]
    if lang == 'UZ':
        lines.insert(7, 'Апостроф (oʻ, gʻ) можно ставить любой, мы приведём его к нужному виду.')
    for index, line in enumerate(lines):
        cell = sheet.cell(row=4 + index, column=2, value=f'{index + 1}. {line}')
        cell.font = Font(name=FONT, size=11)
        sheet.merge_cells(start_row=4 + index, start_column=2, end_row=4 + index, end_column=7)
        cell.alignment = Alignment(wrap_text=True, vertical='center')
        sheet.row_dimensions[4 + index].height = 30
    row = 5 + len(lines)
    sheet.cell(row=row, column=2, value='Пример заполнения (это не настоящая правка):').font = Font(name=FONT, size=11, bold=True)
    example_headers = ['Где в игре', 'По-русски (смысл)', 'Текст сейчас', 'Верно?', 'Как правильно', 'Комментарий']
    example = {
        'UZ': ['Счёт', 'Посчитай! Сколько тут?', 'Sanab ko‘r! Nechta?', 'Нет', 'Sanab ko‘r! Bu yerda nechta?', 'так понятнее детям'],
        'KG': ['Счёт', 'Посчитай! Сколько тут?', 'Санап көр! Канча?', 'Нет', 'Санап көр! Бул жерде канча?', 'так понятнее детям']
    }[lang]
    for offset, (header, value) in enumerate(zip(example_headers, example)):
        head = sheet.cell(row=row + 1, column=2 + offset, value=header)
        head.font = Font(name=FONT, size=10, bold=True, color='FFFFFF')
        head.fill = HEADER_FILL
        head.alignment = CENTER
        cell = sheet.cell(row=row + 2, column=2 + offset, value=value)
        cell.font = Font(name=FONT, size=11)
        cell.fill = FILL_IN if offset >= 3 else EXAMPLE_FILL
        cell.alignment = WRAP
        cell.border = BORDER
    sheet.row_dimensions[row + 2].height = 34
    legend = sheet.cell(row=row + 4, column=2, value='Жёлтые ячейки — для ваших правок. Остальное менять не нужно.')
    legend.font = Font(name=FONT, size=11, italic=True)
    legend.fill = FILL_IN


def write_texts(sheet, lang, items):
    name, _ = LANG_NAMES[lang]
    sheet.title = name
    for index, (header, width) in enumerate(zip(HEADERS, WIDTHS), start=1):
        cell = sheet.cell(row=1, column=index, value=header)
        cell.font = Font(name=FONT, size=11, bold=True, color='FFFFFF')
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = BORDER
        sheet.column_dimensions[cell.column_letter].width = width
    sheet.row_dimensions[1].height = 30
    sheet.freeze_panes = 'A2'
    answers = DataValidation(type='list', formula1='"Да,Нет"', allow_blank=True)
    sheet.add_data_validation(answers)

    for number, item in enumerate(items, start=1):
        row = number + 1
        values = [number, None, item['screen'], item['audience'], context_of(item), item['text'], None, None, None, item['id']]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=row, column=column, value=value)
            cell.font = Font(name=FONT, size=12 if column == 6 else 10, bold=column == 6, color='8A96B8' if column in (1, 10) else '22305A')
            cell.alignment = CENTER if column in (1, 2, 4, 7) else WRAP
            cell.border = BORDER
            if column in (7, 8, 9):
                cell.fill = FILL_IN
        answers.add(sheet.cell(row=row, column=7))
        picture = picture_file(item.get('picture')) if item['screen'] in ('Слова к картинкам', 'Слова по слогам') or item.get('picture') else None
        if picture:
            image = XLImage(thumbnail(picture))
            image.width = image.height = 44
            sheet.add_image(image, f'B{row}')
            sheet.row_dimensions[row].height = 40
        else:
            sheet.row_dimensions[row].height = 30 if len(item['text']) < 60 else 46
    sheet.column_dimensions['J'].hidden = True
    sheet.auto_filter.ref = f'A1:I{len(items) + 1}'


def build(texts_file, out_dir):
    with open(texts_file, encoding='utf-8') as handle:
        items = json.load(handle)
    os.makedirs(out_dir, exist_ok=True)
    for lang in ('UZ', 'KG'):
        rows = [item for item in items if item['lang'] == lang]
        workbook = Workbook()
        write_instructions(workbook.active, lang)
        write_texts(workbook.create_sheet(), lang, rows)
        workbook.active = 0
        name = LANG_NAMES[lang][0].lower()
        target = os.path.join(out_dir, f'Letter Loom — проверка текстов ({name}).xlsx')
        workbook.save(target)
        print(f'{lang}: {len(rows)} texts -> {target}')


def read(filled, out_file):
    workbook = load_workbook(filled)
    corrections, attention = [], []
    for sheet in workbook.worksheets:
        header = [cell.value for cell in sheet[1]]
        if 'id' not in header:
            continue
        col = {name: header.index(name) for name in header if name}
        lang = 'UZ' if sheet.title.startswith('Узбек') else 'KG'
        for row in sheet.iter_rows(min_row=2, values_only=True):
            item_id = row[col['id']]
            if not item_id:
                continue
            original = row[col['Текст сейчас']] or ''
            fix = (row[col['Как правильно']] or '').strip() if isinstance(row[col['Как правильно']], str) else ''
            verdict = (row[col['Верно?']] or '').strip() if isinstance(row[col['Верно?']], str) else ''
            note = row[col['Комментарий']] or ''
            if fix and fix != original:
                corrections.append({'id': item_id, 'lang': lang, 'original': original, 'text': fix, 'comment': note})
            elif verdict.lower().startswith('нет') or note:
                attention.append({'id': item_id, 'lang': lang, 'text': original, 'verdict': verdict, 'comment': note})
    with open(out_file, 'w', encoding='utf-8') as handle:
        json.dump(corrections, handle, ensure_ascii=False, indent=1)
    print(f'{len(corrections)} corrections -> {out_file}')
    for item in attention:
        print(f"  ? {item['lang']} {item['text']!r}: {item['verdict']} {item['comment']}".rstrip())


if __name__ == '__main__':
    if len(sys.argv) == 4 and sys.argv[1] == 'build':
        build(sys.argv[2], sys.argv[3])
    elif len(sys.argv) == 4 and sys.argv[1] == 'read':
        read(sys.argv[2], sys.argv[3])
    else:
        print(__doc__)
        sys.exit(1)
